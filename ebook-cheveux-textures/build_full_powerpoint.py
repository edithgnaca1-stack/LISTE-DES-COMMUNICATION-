# -*- coding: utf-8 -*-
"""
Constructeur Master PowerPoint — Ma couronne, mes règles
Version complète de 75 slides aérées, spacieuses et pédagogiques.
Sortie : MA-COURONNE-MES-REGLES-Complet.pptx
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

from slides_data_intro import SLIDES_INTRO
from slides_data_diag_ingr import SLIDES_DIAG_INGR
from slides_data_routines import SLIDES_ROUTINES
from slides_data_tresses_outils import SLIDES_TRESSES_OUTILS

HERE = os.path.dirname(os.path.abspath(__file__))
IMG_DIR = os.path.join(HERE, "extrait", "images")
OUT_PPTX = os.path.join(HERE, "MA-COURONNE-MES-REGLES-Complet.pptx")

# ------------------------------------------------------------------ COULEURS
C_DARK_BG       = RGBColor(0x1B, 0x27, 0x40) # Indigo nuit
C_LIGHT_BG      = RGBColor(0xFB, 0xF9, 0xF5) # Crème douce lumineuse
C_CARD_BG       = RGBColor(0xFF, 0xFF, 0xFF) # Blanc pur
C_BORDER        = RGBColor(0xE2, 0xD7, 0xC4) # Sable doux
C_TEXT_DARK     = RGBColor(0x1C, 0x1B, 0x19) # Charbon
C_TEXT_MUTED    = RGBColor(0x6E, 0x6A, 0x62) # Gris chaud
C_TEXT_LIGHT    = RGBColor(0xFA, 0xF6, 0xF0) # Blanc cassé

C_GOLD          = RGBColor(0xB8, 0x86, 0x2B) # Or chaud
C_GOLD_LIGHT    = RGBColor(0xD9, 0xA1, 0x3B) # Or éclatant
C_TERRA         = RGBColor(0xA8, 0x52, 0x36) # Terracotta
C_GREEN         = RGBColor(0x55, 0x66, 0x3A) # Vert gombo
C_INDIGO        = RGBColor(0x2B, 0x3E, 0x63) # Indigo primaire

ACCENT_COLORS = {
    "indigo": C_INDIGO,
    "gold": C_GOLD,
    "terra": C_TERRA,
    "green": C_GREEN
}

SLIDE_W = 13.333
SLIDE_H = 7.5

def init_prs():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    return prs

def set_bg(slide, color):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(SLIDE_W), Inches(SLIDE_H))
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    return bg

def add_header(slide, tracker, title, subtitle=None):
    # Tracker
    tb_tr = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.73), Inches(0.3))
    tf_tr = tb_tr.text_frame
    tf_tr.word_wrap = True
    tf_tr.margin_left = tf_tr.margin_top = tf_tr.margin_right = tf_tr.margin_bottom = 0
    p_tr = tf_tr.paragraphs[0]
    p_tr.text = tracker.upper()
    p_tr.font.name = "Arial"
    p_tr.font.size = Pt(9.5)
    p_tr.font.bold = True
    p_tr.font.color.rgb = C_TERRA

    # Titre
    tb_ti = slide.shapes.add_textbox(Inches(0.8), Inches(0.68), Inches(11.73), Inches(0.55))
    tf_ti = tb_ti.text_frame
    tf_ti.word_wrap = True
    tf_ti.margin_left = tf_ti.margin_top = tf_ti.margin_right = tf_ti.margin_bottom = 0
    p_ti = tf_ti.paragraphs[0]
    p_ti.text = title
    p_ti.font.name = "Georgia"
    p_ti.font.size = Pt(20)
    p_ti.font.bold = True
    p_ti.font.color.rgb = C_INDIGO

    # Sous-titre
    if subtitle:
        tb_sub = slide.shapes.add_textbox(Inches(0.8), Inches(1.22), Inches(11.73), Inches(0.35))
        tf_sub = tb_sub.text_frame
        tf_sub.word_wrap = True
        tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = subtitle
        p_sub.font.name = "Arial"
        p_sub.font.size = Pt(11)
        p_sub.font.color.rgb = C_TEXT_MUTED

def add_footer(slide, slide_num, total_slides):
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(7.08), Inches(11.733), Inches(0.012))
    line.fill.solid()
    line.fill.fore_color.rgb = C_BORDER
    line.line.fill.background()

    tb = slide.shapes.add_textbox(Inches(0.8), Inches(7.13), Inches(8.5), Inches(0.28))
    tf = tb.text_frame
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = "Ma couronne, mes règles — Le guide complet des cheveux texturés en Afrique  ·  Édition Bénin 2026"
    p.font.name = "Arial"
    p.font.size = Pt(8.5)
    p.font.color.rgb = C_TEXT_MUTED

    tb2 = slide.shapes.add_textbox(Inches(9.5), Inches(7.13), Inches(3.033), Inches(0.28))
    tf2 = tb2.text_frame
    tf2.margin_left = tf2.margin_top = tf2.margin_right = tf2.margin_bottom = 0
    p2 = tf2.paragraphs[0]
    p2.alignment = PP_ALIGN.RIGHT
    p2.text = f"{slide_num} / {total_slides}"
    p2.font.name = "Arial"
    p2.font.size = Pt(8.5)
    p2.font.bold = True
    p2.font.color.rgb = C_GOLD

def add_callout(slide, text, top=5.95, label="À RETENIR"):
    # Nettoyage si le texte commence déjà par le label
    clean_text = text.strip()
    for prefix in ["À RETENIR —", "À RETENIR :", "À RETENIR", "A RETENIR —", "A RETENIR :"]:
        if clean_text.startswith(prefix):
            clean_text = clean_text[len(prefix):].strip()
            if clean_text.startswith("—") or clean_text.startswith(":"):
                clean_text = clean_text[1:].strip()

    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(top), Inches(11.733), Inches(0.95))
    box.fill.solid()
    box.fill.fore_color.rgb = C_DARK_BG
    box.line.fill.background()

    tb = slide.shapes.add_textbox(Inches(1.1), Inches(top + 0.15), Inches(11.133), Inches(0.65))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = f"{label} — "
    p.font.name = "Arial"
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = C_GOLD_LIGHT

    run = p.add_run()
    run.text = clean_text
    run.font.name = "Arial"
    run.font.size = Pt(10.5)
    run.font.bold = False
    run.font.color.rgb = C_TEXT_LIGHT

def create_cover(prs, data, slide_num, total_slides):
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    set_bg(slide, C_DARK_BG)

    # Cadre doré fin autour de la slide
    border = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(0.4), Inches(12.533), Inches(6.7))
    border.fill.background()
    border.line.color.rgb = C_GOLD
    border.line.width = Pt(1.5)

    # Bandeau supérieur
    tb_top = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(7.5), Inches(0.35))
    tf_top = tb_top.text_frame
    tf_top.margin_left = tf_top.margin_top = tf_top.margin_right = tf_top.margin_bottom = 0
    p_top = tf_top.paragraphs[0]
    p_top.text = "GUIDE PRATIQUE COMPLET  ·  CHEVEUX TEXTURÉS  ·  BÉNIN & AFRIQUE"
    p_top.font.name = "Arial"
    p_top.font.size = Pt(10)
    p_top.font.bold = True
    p_top.font.color.rgb = C_GOLD_LIGHT

    # Titre majeur
    tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(1.15), Inches(7.5), Inches(1.6))
    tf_title = tb_title.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
    p_t1 = tf_title.paragraphs[0]
    p_t1.text = "Ma couronne,\nmes règles"
    p_t1.font.name = "Georgia"
    p_t1.font.size = Pt(40)
    p_t1.font.bold = True
    p_t1.font.color.rgb = C_TEXT_LIGHT

    # Sous-titre
    tb_sub = slide.shapes.add_textbox(Inches(0.8), Inches(2.9), Inches(7.2), Inches(0.8))
    tf_sub = tb_sub.text_frame
    tf_sub.word_wrap = True
    tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0
    p_s = tf_sub.paragraphs[0]
    p_s.text = data["tagline"]
    p_s.font.name = "Arial"
    p_s.font.size = Pt(13)
    p_s.font.color.rgb = RGBColor(0xD2, 0xC8, 0xB8)

    # 4 Cartouches des promesses (2x2)
    promesses = [
        ("POURQUOI ÇA MARCHE ?", "Ce que chaque ingrédient fait à la fibre, expliqué simplement."),
        ("COMBIEN EXACTEMENT ?", "Cuillères, verres, noix : des quantités précises selon la longueur."),
        ("À QUELLE FRÉQUENCE ?", "Chaque geste a son rythme : semaine, mois ou saison."),
        ("LE KIT ANTI-PARESSE", "Le minimum vital qui marche quand vous n'avez ni le temps ni l'envie.")
    ]
    for idx, (p_tit, p_txt) in enumerate(promesses):
        row = idx // 2
        col = idx % 2
        bx = 0.8 + col * 3.65
        by = 3.85 + row * 1.35
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(bx), Inches(by), Inches(3.5), Inches(1.2))
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(0x24, 0x33, 0x52)
        card.line.color.rgb = C_GOLD
        card.line.width = Pt(0.7)

        tb_c = slide.shapes.add_textbox(Inches(bx + 0.15), Inches(by + 0.12), Inches(3.2), Inches(0.95))
        tf_c = tb_c.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = tf_c.margin_top = tf_c.margin_right = tf_c.margin_bottom = 0
        p1 = tf_c.paragraphs[0]
        p1.text = p_tit
        p1.font.name = "Arial"
        p1.font.size = Pt(10)
        p1.font.bold = True
        p1.font.color.rgb = C_GOLD_LIGHT
        p1.space_after = Pt(3)

        p2 = tf_c.add_paragraph()
        p2.text = p_txt
        p2.font.name = "Arial"
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = C_TEXT_LIGHT

    # Image de couverture sur la droite
    img_path = os.path.join(IMG_DIR, data.get("image", "cover_crop.png"))
    if os.path.exists(img_path):
        slide.shapes.add_picture(img_path, Inches(8.5), Inches(1.0), Inches(4.0), Inches(5.33))
        # Cadre autour de l'image
        f_pic = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.5), Inches(1.0), Inches(4.0), Inches(5.33))
        f_pic.fill.background()
        f_pic.line.color.rgb = C_GOLD
        f_pic.line.width = Pt(1.5)

    # Pied de couverture
    tb_foot = slide.shapes.add_textbox(Inches(0.8), Inches(6.55), Inches(7.5), Inches(0.35))
    tf_foot = tb_foot.text_frame
    tf_foot.margin_left = tf_foot.margin_top = tf_foot.margin_right = tf_foot.margin_bottom = 0
    p_f = tf_foot.paragraphs[0]
    p_f.text = f"{data['author']}  ·  {data['date']}  ·  Document de travail complet"
    p_f.font.name = "Arial"
    p_f.font.size = Pt(9.5)
    p_f.font.color.rgb = C_GOLD

def create_divider(prs, data, slide_num, total_slides):
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    set_bg(slide, C_DARK_BG)

    # Cadre intérieur
    border = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(0.8), Inches(11.333), Inches(5.9))
    border.fill.background()
    border.line.color.rgb = C_GOLD
    border.line.width = Pt(1.2)

    # Numéro de partie
    tb_num = slide.shapes.add_textbox(Inches(1.5), Inches(1.3), Inches(10.33), Inches(0.4))
    tf_num = tb_num.text_frame
    tf_num.margin_left = tf_num.margin_top = tf_num.margin_right = tf_num.margin_bottom = 0
    p_num = tf_num.paragraphs[0]
    p_num.text = data["part_num"].upper()
    p_num.font.name = "Arial"
    p_num.font.size = Pt(12)
    p_num.font.bold = True
    p_num.font.color.rgb = C_GOLD_LIGHT

    # Titre de partie
    tb_tit = slide.shapes.add_textbox(Inches(1.5), Inches(1.8), Inches(10.33), Inches(1.2))
    tf_tit = tb_tit.text_frame
    tf_tit.word_wrap = True
    tf_tit.margin_left = tf_tit.margin_top = tf_tit.margin_right = tf_tit.margin_bottom = 0
    p_tit = tf_tit.paragraphs[0]
    p_tit.text = data["title"]
    p_tit.font.name = "Georgia"
    p_tit.font.size = Pt(36)
    p_tit.font.bold = True
    p_tit.font.color.rgb = C_TEXT_LIGHT

    # Ligne décorative dorée
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.5), Inches(3.2), Inches(3.5), Inches(0.04))
    line.fill.solid()
    line.fill.fore_color.rgb = C_GOLD_LIGHT
    line.line.fill.background()

    # Citation
    tb_q = slide.shapes.add_textbox(Inches(1.5), Inches(3.45), Inches(10.33), Inches(1.4))
    tf_q = tb_q.text_frame
    tf_q.word_wrap = True
    tf_q.margin_left = tf_q.margin_top = tf_q.margin_right = tf_q.margin_bottom = 0
    p_q = tf_q.paragraphs[0]
    p_q.text = data["quote"]
    p_q.font.name = "Georgia"
    p_q.font.size = Pt(18)
    p_q.font.italic = True
    p_q.font.color.rgb = C_GOLD_LIGHT

    # Subtitle / lede
    tb_sub = slide.shapes.add_textbox(Inches(1.5), Inches(5.0), Inches(10.33), Inches(1.0))
    tf_sub = tb_sub.text_frame
    tf_sub.word_wrap = True
    tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0
    p_s = tf_sub.paragraphs[0]
    p_s.text = data["subtitle"]
    p_s.font.name = "Arial"
    p_s.font.size = Pt(13)
    p_s.font.color.rgb = RGBColor(0xD2, 0xC8, 0xB8)

def render_card_contents(tf, card_data, accent_color):
    p0 = tf.paragraphs[0]
    p0.text = card_data["title"]
    p0.font.name = "Georgia"
    p0.font.size = Pt(14.5)
    p0.font.bold = True
    p0.font.color.rgb = accent_color
    p0.space_after = Pt(8)

    # Paragraphs (tuples de (titre, corps) ou chaînes simples)
    if "paragraphs" in card_data:
        for idx, item in enumerate(card_data["paragraphs"]):
            p = tf.add_paragraph()
            if isinstance(item, tuple) and len(item) == 2:
                p_tit, p_body = item
                p.text = p_tit + " "
                p.font.name = "Arial"
                p.font.size = Pt(11)
                p.font.bold = True
                p.font.color.rgb = accent_color
                p.space_after = Pt(4)

                run = p.add_run()
                run.text = p_body
                run.font.name = "Arial"
                run.font.size = Pt(11)
                run.font.bold = False
                run.font.color.rgb = C_TEXT_DARK
            else:
                p.text = "• " + str(item)
                p.font.name = "Arial"
                p.font.size = Pt(11)
                p.font.color.rgb = C_TEXT_DARK
                p.space_after = Pt(4)

    # Body (liste de bullets simples)
    elif "body" in card_data:
        for item in card_data["body"]:
            p = tf.add_paragraph()
            p.text = "• " + item
            p.font.name = "Arial"
            p.font.size = Pt(11)
            p.font.color.rgb = C_TEXT_DARK
            p.space_after = Pt(4)

def create_two_card(prs, data, slide_num, total_slides):
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    set_bg(slide, C_LIGHT_BG)
    add_header(slide, data["tracker"], data["title"], data.get("subtitle"))

    card_h = 4.25 if "callout" in data else 5.25

    # Carte 1 (Gauche)
    c1 = data["card1"]
    acc1 = ACCENT_COLORS.get(c1.get("accent", "indigo"), C_INDIGO)
    card_shape1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.65), Inches(card_h))
    card_shape1.fill.solid()
    card_shape1.fill.fore_color.rgb = C_CARD_BG
    card_shape1.line.color.rgb = C_BORDER
    card_shape1.line.width = Pt(1)

    bar1 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.95), Inches(1.62), Inches(5.35), Inches(0.05))
    bar1.fill.solid()
    bar1.fill.fore_color.rgb = acc1
    bar1.line.fill.background()

    tb1 = slide.shapes.add_textbox(Inches(1.05), Inches(1.8), Inches(5.15), Inches(card_h - 0.35))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = tf1.margin_top = tf1.margin_right = tf1.margin_bottom = 0
    render_card_contents(tf1, c1, acc1)

    # Carte 2 (Droite)
    c2 = data["card2"]
    acc2 = ACCENT_COLORS.get(c2.get("accent", "terra"), C_TERRA)
    card_shape2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.6), Inches(5.65), Inches(card_h))
    card_shape2.fill.solid()
    card_shape2.fill.fore_color.rgb = C_CARD_BG
    card_shape2.line.color.rgb = C_BORDER
    card_shape2.line.width = Pt(1)

    bar2 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.0), Inches(1.62), Inches(5.35), Inches(0.05))
    bar2.fill.solid()
    bar2.fill.fore_color.rgb = acc2
    bar2.line.fill.background()

    tb2 = slide.shapes.add_textbox(Inches(7.1), Inches(1.8), Inches(5.15), Inches(card_h - 0.35))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    tf2.margin_left = tf2.margin_top = tf2.margin_right = tf2.margin_bottom = 0
    render_card_contents(tf2, c2, acc2)

    if "callout" in data:
        add_callout(slide, data["callout"])

    add_footer(slide, slide_num, total_slides)

def create_three_card(prs, data, slide_num, total_slides):
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    set_bg(slide, C_LIGHT_BG)
    add_header(slide, data["tracker"], data["title"], data.get("subtitle"))

    card_h = 4.25 if "callout" in data else 5.25
    card_w = 3.65
    gap = 0.39

    cols = [data["col1"], data["col2"], data["col3"]]
    for idx, c in enumerate(cols):
        left = 0.8 + idx * (card_w + gap)
        acc = ACCENT_COLORS.get(c.get("accent", "indigo"), C_INDIGO)

        card_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(1.6), Inches(card_w), Inches(card_h))
        card_shape.fill.solid()
        card_shape.fill.fore_color.rgb = C_CARD_BG
        card_shape.line.color.rgb = C_BORDER
        card_shape.line.width = Pt(1)

        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left + 0.15), Inches(1.62), Inches(card_w - 0.3), Inches(0.05))
        bar.fill.solid()
        bar.fill.fore_color.rgb = acc
        bar.line.fill.background()

        tb = slide.shapes.add_textbox(Inches(left + 0.2), Inches(1.8), Inches(card_w - 0.4), Inches(card_h - 0.35))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        render_card_contents(tf, c, acc)

    if "callout" in data:
        add_callout(slide, data["callout"])

    add_footer(slide, slide_num, total_slides)

def create_four_boxes(prs, data, slide_num, total_slides):
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    set_bg(slide, C_LIGHT_BG)
    add_header(slide, data["tracker"], data["title"], data.get("subtitle"))

    boxes = [data["box1"], data["box2"], data["box3"], data["box4"]]
    box_w = 5.65
    box_h = 2.45
    coords = [
        (0.8, 1.65), (6.85, 1.65),
        (0.8, 4.45), (6.85, 4.45)
    ]
    colors = [C_TERRA, C_INDIGO, C_GREEN, C_GOLD]

    for idx, (tit, sub, txt) in enumerate(boxes):
        bx, by = coords[idx]
        col = colors[idx % len(colors)]

        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(bx), Inches(by), Inches(box_w), Inches(box_h))
        card.fill.solid()
        card.fill.fore_color.rgb = C_CARD_BG
        card.line.color.rgb = C_BORDER
        card.line.width = Pt(1)

        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(bx + 0.15), Inches(by + 0.02), Inches(box_w - 0.3), Inches(0.04))
        bar.fill.solid()
        bar.fill.fore_color.rgb = col
        bar.line.fill.background()

        tb = slide.shapes.add_textbox(Inches(bx + 0.2), Inches(by + 0.15), Inches(box_w - 0.4), Inches(box_h - 0.25))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = tit
        p0.font.name = "Georgia"
        p0.font.size = Pt(13)
        p0.font.bold = True
        p0.font.color.rgb = col

        p1 = tf.add_paragraph()
        p1.text = sub.upper()
        p1.font.name = "Arial"
        p1.font.size = Pt(9.5)
        p1.font.bold = True
        p1.font.color.rgb = C_TEXT_MUTED
        p1.space_after = Pt(4)

        p2 = tf.add_paragraph()
        p2.text = txt
        p2.font.name = "Arial"
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = C_TEXT_DARK

    add_footer(slide, slide_num, total_slides)

def create_table_slide(prs, data, slide_num, total_slides):
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    set_bg(slide, C_LIGHT_BG)
    add_header(slide, data["tracker"], data["title"], data.get("subtitle"))

    headers = data["headers"]
    rows = data["rows"]
    num_rows = len(rows) + 1
    num_cols = len(headers)

    table_shape = slide.shapes.add_table(num_rows, num_cols, Inches(0.8), Inches(1.6), Inches(11.733), Inches(4.2))
    table = table_shape.table

    # Format Header
    for c_idx, h_text in enumerate(headers):
        cell = table.cell(0, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = C_DARK_BG
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = h_text
        p.font.name = "Arial"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = C_GOLD_LIGHT

    # Format Rows
    for r_idx, row_data in enumerate(rows):
        bg_col = C_CARD_BG if r_idx % 2 == 0 else RGBColor(0xF5, 0xEE, 0xE0)
        for c_idx, cell_value in enumerate(row_data):
            cell = table.cell(r_idx + 1, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg_col
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = cell_value
            p.font.name = "Arial"
            p.font.size = Pt(9.5)
            if c_idx == 0:
                p.font.bold = True
                p.font.color.rgb = C_INDIGO
            else:
                p.font.color.rgb = C_TEXT_DARK

    if "callout" in data:
        add_callout(slide, data["callout"], top=5.95)

    add_footer(slide, slide_num, total_slides)

def create_image_card(prs, data, slide_num, total_slides):
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    set_bg(slide, C_LIGHT_BG)
    add_header(slide, data["tracker"], data["title"], data.get("subtitle"))

    # Image sur la gauche
    img_path = os.path.join(IMG_DIR, data.get("image", "photo_mesures.png"))
    if os.path.exists(img_path):
        slide.shapes.add_picture(img_path, Inches(0.8), Inches(1.6), Inches(4.6), Inches(4.15))
        f_pic = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.6), Inches(4.6), Inches(4.15))
        f_pic.fill.background()
        f_pic.line.color.rgb = C_BORDER
        f_pic.line.width = Pt(1)

    # Carte sur la droite
    c = data["card"]
    acc = ACCENT_COLORS.get(c.get("accent", "gold"), C_GOLD)
    card_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.7), Inches(1.6), Inches(6.833), Inches(4.15))
    card_shape.fill.solid()
    card_shape.fill.fore_color.rgb = C_CARD_BG
    card_shape.line.color.rgb = C_BORDER
    card_shape.line.width = Pt(1)

    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5.85), Inches(1.62), Inches(6.533), Inches(0.05))
    bar.fill.solid()
    bar.fill.fore_color.rgb = acc
    bar.line.fill.background()

    tb = slide.shapes.add_textbox(Inches(5.95), Inches(1.8), Inches(6.333), Inches(3.8))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    render_card_contents(tf, c, acc)

    if "callout" in data:
        add_callout(slide, data["callout"], top=5.95)

    add_footer(slide, slide_num, total_slides)

def build_presentation():
    all_slides_data = []
    all_slides_data.extend(SLIDES_INTRO)
    all_slides_data.extend(SLIDES_DIAG_INGR)
    all_slides_data.extend(SLIDES_ROUTINES)
    all_slides_data.extend(SLIDES_TRESSES_OUTILS)

    total_slides = len(all_slides_data)
    print(f"Compilation de {total_slides} slides PowerPoint...")

    prs = init_prs()

    for idx, slide_data in enumerate(all_slides_data):
        slide_num = idx + 1
        s_type = slide_data.get("type", "two_card")

        if s_type == "cover":
            create_cover(prs, slide_data, slide_num, total_slides)
        elif s_type == "divider":
            create_divider(prs, slide_data, slide_num, total_slides)
        elif s_type == "two_card":
            create_two_card(prs, slide_data, slide_num, total_slides)
        elif s_type == "three_card":
            create_three_card(prs, slide_data, slide_num, total_slides)
        elif s_type == "four_boxes":
            create_four_boxes(prs, slide_data, slide_num, total_slides)
        elif s_type == "table":
            create_table_slide(prs, slide_data, slide_num, total_slides)
        elif s_type == "image_card":
            create_image_card(prs, slide_data, slide_num, total_slides)
        else:
            create_two_card(prs, slide_data, slide_num, total_slides)

    prs.save(OUT_PPTX)
    print(f"Succès ! Présentation enregistrée : {OUT_PPTX}")
    print(f"Nombre total de slides : {len(prs.slides)}")
    return OUT_PPTX

if __name__ == "__main__":
    build_presentation()
