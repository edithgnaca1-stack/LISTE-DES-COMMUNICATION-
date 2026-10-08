# -*- coding: utf-8 -*-
"""
Moteur de génération PowerPoint pour l'e-book :
Ma couronne, mes règles — Le guide complet des cheveux texturés en Afrique
Design : 16:9 widescreen, aéré, élégant, lisibilité maximale.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import os

# ------------------------------------------------------------------ COULEURS
C_DARK_BG       = RGBColor(0x1B, 0x27, 0x40) # Indigo nuit profond
C_LIGHT_BG      = RGBColor(0xFB, 0xF9, 0xF5) # Crème très douce et lumineuse
C_CARD_BG       = RGBColor(0xFF, 0xFF, 0xFF) # Blanc pur pour cartes
C_CARD_WARM     = RGBColor(0xF5, 0xEE, 0xE0) # Karité chaud doux
C_BORDER        = RGBColor(0xE2, 0xD7, 0xC4) # Sable clair
C_TEXT_DARK     = RGBColor(0x1C, 0x1B, 0x19) # Encre charbon
C_TEXT_MUTED    = RGBColor(0x6E, 0x6A, 0x62) # Gris chaud
C_TEXT_LIGHT    = RGBColor(0xFA, 0xF6, 0xF0) # Blanc cassé

C_GOLD          = RGBColor(0xB8, 0x86, 0x2B) # Or chaud noble
C_GOLD_LIGHT    = RGBColor(0xD9, 0xA1, 0x3B) # Or éclatant
C_TERRA         = RGBColor(0xA8, 0x52, 0x36) # Terracotta riche
C_GREEN         = RGBColor(0x55, 0x66, 0x3A) # Vert gombo végétal
C_INDIGO        = RGBColor(0x2B, 0x3E, 0x63) # Indigo primaire

# ------------------------------------------------------------------ DIMENSIONS
SLIDE_W = 13.333
SLIDE_H = 7.5

def init_presentation():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    return prs

def set_slide_background(slide, color):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(SLIDE_W), Inches(SLIDE_H))
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    return bg

def add_header(slide, tracker, title, subtitle=None):
    """Bandeau d'en-tête standard et aéré pour slide de contenu"""
    # Tracker (petite catégorie en haut)
    tb_tr = slide.shapes.add_textbox(Inches(0.8), Inches(0.42), Inches(11.73), Inches(0.3))
    tf_tr = tb_tr.text_frame
    tf_tr.word_wrap = True
    tf_tr.margin_left = tf_tr.margin_top = tf_tr.margin_right = tf_tr.margin_bottom = 0
    p_tr = tf_tr.paragraphs[0]
    p_tr.text = tracker.upper()
    p_tr.font.name = "Arial"
    p_tr.font.size = Pt(9.5)
    p_tr.font.bold = True
    p_tr.font.color.rgb = C_TERRA

    # Titre principal
    tb_ti = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.73), Inches(0.55))
    tf_ti = tb_ti.text_frame
    tf_ti.word_wrap = True
    tf_ti.margin_left = tf_ti.margin_top = tf_ti.margin_right = tf_ti.margin_bottom = 0
    p_ti = tf_ti.paragraphs[0]
    p_ti.text = title
    p_ti.font.name = "Georgia"
    p_ti.font.size = Pt(21)
    p_ti.font.bold = True
    p_ti.font.color.rgb = C_INDIGO

    # Sous-titre optionnel
    if subtitle:
        tb_sub = slide.shapes.add_textbox(Inches(0.8), Inches(1.25), Inches(11.73), Inches(0.35))
        tf_sub = tb_sub.text_frame
        tf_sub.word_wrap = True
        tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = subtitle
        p_sub.font.name = "Arial"
        p_sub.font.size = Pt(11.5)
        p_sub.font.color.rgb = C_TEXT_MUTED

def add_footer(slide, slide_num, total_slides=None):
    """Pied de page discret et élégant"""
    # Ligne fine
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(7.08), Inches(11.733), Inches(0.015))
    line.fill.solid()
    line.fill.fore_color.rgb = C_BORDER
    line.line.fill.background()

    # Texte gauche
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(7.13), Inches(8.0), Inches(0.28))
    tf = tb.text_frame
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = "Ma couronne, mes règles — Le guide complet des cheveux texturés en Afrique  ·  Édition 2026"
    p.font.name = "Arial"
    p.font.size = Pt(8.5)
    p.font.color.rgb = C_TEXT_MUTED

    # Texte droite (numéro)
    tb2 = slide.shapes.add_textbox(Inches(9.5), Inches(7.13), Inches(3.033), Inches(0.28))
    tf2 = tb2.text_frame
    tf2.margin_left = tf2.margin_top = tf2.margin_right = tf2.margin_bottom = 0
    p2 = tf2.paragraphs[0]
    p2.alignment = PP_ALIGN.RIGHT
    p2.text = f"Page {slide_num}" if not total_slides else f"{slide_num} / {total_slides}"
    p2.font.name = "Arial"
    p2.font.size = Pt(8.5)
    p2.font.bold = True
    p2.font.color.rgb = C_GOLD

def add_card(slide, left, top, width, height, title=None, accent_color=C_INDIGO, bg_color=C_CARD_BG):
    """Crée une carte propre avec fond blanc, bordure douce et barre d'accent"""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = C_BORDER
    card.line.width = Pt(1)

    # Petite barre d'accent en haut de la carte
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left + 0.15), Inches(top + 0.02), Inches(width - 0.3), Inches(0.05))
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent_color
    bar.line.fill.background()

    # Zone de texte
    top_offset = top + 0.2
    h_offset = height - 0.3
    tb = slide.shapes.add_textbox(Inches(left + 0.25), Inches(top_offset), Inches(width - 0.5), Inches(h_offset))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    if title:
        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.name = "Georgia"
        p0.font.size = Pt(14)
        p0.font.bold = True
        p0.font.color.rgb = accent_color
        p0.space_after = Pt(8)

    return tf

def add_callout_banner(slide, text, left=0.8, top=5.9, width=11.733, height=0.95, label="À RETENIR", color=C_INDIGO):
    """Bandeau de synthèse fort en bas de slide"""
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    box.fill.solid()
    box.fill.fore_color.rgb = color
    box.line.fill.background()

    tb = slide.shapes.add_textbox(Inches(left + 0.3), Inches(top + 0.15), Inches(width - 0.6), Inches(height - 0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = f"{label} — "
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_GOLD_LIGHT

    run = p.add_run()
    run.text = text
    run.font.name = "Arial"
    run.font.size = Pt(11)
    run.font.bold = False
    run.font.color.rgb = C_TEXT_LIGHT

print("generator_engine.py défini avec succès")
